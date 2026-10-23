import argparse
import statistics
import sys
from pathlib import Path
import pandas as pd

DEAL = {
    "offer_price": 31.00,          # Merger agreement, 27 Feb 2026 (WBD 8-K)
    "unaffected_price": 12.54,     # Close on 10 Sep 2025, before reports of Paramount's interest (DEFM14A)
    "equity_value": 81_000,        # Paramount announcement: ~$81bn equity value
    "enterprise_value": 110_000,   # Paramount announcement: $110bn EV
    "synergies_run_rate": 6_000,   # Paramount: "more than $6 billion" of synergies
    "netflix_offer_cash": 27.75,   # Netflix deal per-share consideration (plus Discovery Global spin shares)
}

WBD = {
    # FY2025, from WBD's Q4 2025 earnings release / 10-K segment note
    "rev_streaming_studios": 20_371,
    "rev_linear": 17_656,
    "ebitda_streaming_studios": 3_496,   # Streaming 1,370 + Studios 2,545 - eliminations 419
    "ebitda_linear": 6_412,              # Global Linear Networks
    "ebitda_corporate": -1_096,          # Corporate costs (allocated pro rata in SOTP)
    "ebitda_other_elims": -68,
    "ebitda_total": 8_744,               # Total adjusted EBITDA
    "net_debt": 29_000,                  # Net debt at 31 Dec 2025
    "diluted_shares": DEAL["equity_value"] / DEAL["offer_price"],
}

ADVISER_RANGES = {
    "Advisers: trading comps": (15.53, 21.18),
    "Advisers: DCF (1)": (21.16, 31.91),
    "Advisers: DCF (2)": (20.67, 31.34),
}

PEERS = {
    "NFLX": ("Netflix", "studio_streaming"),
    "DIS": ("Disney", "studio_streaming"),
    "CMCSA": ("Comcast", "diversified"),
    "FOXA": ("Fox Corp", "linear"),
    "AMCX": ("AMC Networks", "linear"),
    "VSNT": ("Versant Media", "linear"),
}


def fetch_live(tickers):
    """Pull trailing figures from Yahoo Finance. Returns a DataFrame in USD millions."""
    try:
        import yfinance as yf
    except ImportError:
        sys.exit("yfinance is not installed: pip install yfinance")

    rows = []
    for tk in tickers:
        try:
            info = yf.Ticker(tk).info
        except Exception as exc: 
            print(f"  ! {tk}: could not fetch ({exc}); skipping")
            continue
        def m(key):
            v = info.get(key)
            return v / 1e6 if isinstance(v, (int, float)) else None
        rows.append({
            "ticker": tk,
            "price": info.get("currentPrice") or info.get("regularMarketPrice"),
            "market_cap": m("marketCap"),
            "enterprise_value": m("enterpriseValue"),
            "revenue": m("totalRevenue"),
            "ebitda": m("ebitda"),
            "net_income": m("netIncomeToCommon"),
            "revenue_growth": info.get("revenueGrowth"),
            "ebitda_margin": info.get("ebitdaMargins"),
        })
        print(f"  ✓ {tk}")
    return pd.DataFrame(rows)


def build_comps(df):
    df = df.copy()
    df["name"] = df["ticker"].map(lambda t: PEERS.get(t, (t, "other"))[0])
    df["group"] = df["ticker"].map(lambda t: PEERS.get(t, (t, "other"))[1])

    def ratio(a, b):
        return a / b if pd.notna(a) and pd.notna(b) and b > 0 else None

    df["ev_ebitda"] = [ratio(e, x) for e, x in zip(df["enterprise_value"], df["ebitda"])]
    df["ev_revenue"] = [ratio(e, r) for e, r in zip(df["enterprise_value"], df["revenue"])]
    df["pe"] = [ratio(m, n) for m, n in zip(df["market_cap"], df["net_income"])]

    df.loc[(df["ev_ebitda"] <= 0) | (df["ev_ebitda"] > 60), "ev_ebitda"] = None
    return df[["ticker", "name", "group", "price", "market_cap", "enterprise_value",
               "revenue", "ebitda", "ev_ebitda", "ev_revenue", "pe",
               "revenue_growth", "ebitda_margin"]]


def quartiles(values):
    vals = sorted(v for v in values if v is not None and pd.notna(v))
    if not vals:
        return None
    if len(vals) == 1:
        return vals[0], vals[0], vals[0]
    q = statistics.quantiles(vals, n=4, method="inclusive")
    return q[0], statistics.median(vals), q[2]


def per_share(ev):
    return (ev - WBD["net_debt"]) / WBD["diluted_shares"]


def value_wbd(comps):
    results = {}

    q = quartiles(comps["ev_ebitda"])
    if q:
        results["Whole-company comps"] = tuple(per_share(m * WBD["ebitda_total"]) for m in (q[0], q[2]))
        results["_whole_median_mult"] = q[1]

    ss = comps.loc[comps["group"].isin(["studio_streaming", "diversified"]), "ev_ebitda"]
    lin = comps.loc[comps["group"] == "linear", "ev_ebitda"]
    q_ss, q_lin = quartiles(ss), quartiles(lin)
    if q_ss and q_lin:
        overhead = WBD["ebitda_corporate"] + WBD["ebitda_other_elims"]
        gross = WBD["ebitda_streaming_studios"] + WBD["ebitda_linear"]
        ss_net = WBD["ebitda_streaming_studios"] * (1 + overhead / gross)
        lin_net = WBD["ebitda_linear"] * (1 + overhead / gross)
        low = per_share(q_ss[0] * ss_net + q_lin[0] * lin_net)
        high = per_share(q_ss[2] * ss_net + q_lin[2] * lin_net)
        results["Sum-of-the-parts comps"] = (low, high)
        results["_sotp_detail"] = {
            "S&S EBITDA after overhead": ss_net, "S&S multiple range": (q_ss[0], q_ss[2]),
            "Linear EBITDA after overhead": lin_net, "Linear multiple range": (q_lin[0], q_lin[2]),
        }
    return results


def deal_metrics(whole_median_mult):
    ev = DEAL["enterprise_value"]
    shares = WBD["diluted_shares"]
    premium_value = (DEAL["offer_price"] - DEAL["unaffected_price"]) * shares
    m = {
        "Premium to unaffected price": DEAL["offer_price"] / DEAL["unaffected_price"] - 1,
        "EV / FY2025 adjusted EBITDA (pre-synergy)": ev / WBD["ebitda_total"],
        "EV / (FY2025 EBITDA + full synergies)": ev / (WBD["ebitda_total"] + DEAL["synergies_run_rate"]),
        "EV / FY2025 revenue": ev / (WBD["rev_streaming_studios"] + WBD["rev_linear"]),
        "Premium paid ($m)": premium_value,
    }
    if whole_median_mult:
        cap = DEAL["synergies_run_rate"] * whole_median_mult
        m["Synergies capitalised at peer median multiple ($m)"] = cap
        m["Share of synergy value handed to WBD holders"] = premium_value / cap
    return m


def football_field(ranges, out_png, as_of):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = list(ranges.keys())
    fig, ax = plt.subplots(figsize=(9, 0.6 * len(labels) + 1.8))
    for i, lab in enumerate(labels):
        lo, hi = ranges[lab]
        colour = "#8a94a6" if lab.startswith("Advisers") else "#2f6db5"
        ax.barh(i, hi - lo, left=lo, height=0.5, color=colour)
        ax.text(hi + 0.4, i, f"\\${lo:.2f} – \\${hi:.2f}", va="center", fontsize=9)
    ax.set_yticks(range(len(labels)), labels)
    ax.invert_yaxis()
    top = -0.75  # above the first bar once the axis is inverted
    ax.set_ylim(len(labels) - 0.5, top - 0.35)
    ax.axvline(DEAL["offer_price"], color="#c0392b", lw=2)
    ax.text(DEAL["offer_price"], top, " Offer \\$31.00", color="#c0392b", fontsize=9, va="center")
    ax.axvline(DEAL["unaffected_price"], color="#555", lw=1, ls="--")
    ax.text(DEAL["unaffected_price"], top, " Unaffected \\$12.54", color="#555", fontsize=9, va="center")
    ax.set_xlabel("Implied value per WBD share ($)")
    ax.set_title(f"WBD valuation vs $31 offer  (peer data as of {as_of})", fontsize=11)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(out_png, dpi=160)
    plt.close(fig)


def write_excel(comps, ranges, metrics, sotp, out_xlsx, as_of):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.drawing.image import Image as XLImage

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as xw:
        view = comps.rename(columns={
            "market_cap": "Mkt cap ($m)", "enterprise_value": "EV ($m)", "revenue": "Revenue ($m)",
            "ebitda": "EBITDA ($m)", "ev_ebitda": "EV/EBITDA", "ev_revenue": "EV/Revenue", "pe": "P/E",
            "revenue_growth": "Rev growth", "ebitda_margin": "EBITDA margin"})
        view.to_excel(xw, sheet_name="Peer comps", index=False, startrow=1)
        stats = []
        for label, fn in (("25th pct", 0), ("Median", 1), ("75th pct", 2)):
            row = {"name": label}
            for col in ("EV/EBITDA", "EV/Revenue", "P/E"):
                q = quartiles(view[col])
                row[col] = q[fn] if q else None
            stats.append(row)
        pd.DataFrame(stats).to_excel(xw, sheet_name="Peer comps", index=False,
                                     startrow=len(view) + 4, startcol=1)

        pd.DataFrame(
            [(k, v) for k, v in {**DEAL, **WBD}.items()], columns=["Input", "Value"]
        ).to_excel(xw, sheet_name="WBD inputs", index=False)

        val = pd.DataFrame([(k, lo, hi) for k, (lo, hi) in ranges.items()],
                           columns=["Method", "Low ($/sh)", "High ($/sh)"])
        val["Offer vs high"] = DEAL["offer_price"] / val["High ($/sh)"] - 1
        val.to_excel(xw, sheet_name="Valuation", index=False)
        met = pd.DataFrame(list(metrics.items()), columns=["Deal metric", "Value"])
        met.to_excel(xw, sheet_name="Valuation", index=False, startrow=len(val) + 3)
        if sotp:
            pd.DataFrame([(k, f"{v[0]:.1f}x – {v[1]:.1f}x" if isinstance(v, tuple) else v)
                          for k, v in sotp.items()],
                         columns=["SOTP input", "Value"]).to_excel(
                xw, sheet_name="Valuation", index=False, startrow=len(val) + len(met) + 6)

        bold = Font(bold=True)
        head = PatternFill("solid", fgColor="DCE6F1")
        ws = xw.sheets["Peer comps"]
        ws["A1"] = f"Peer trading comps (Yahoo Finance, as of {as_of}; USD millions)"
        ws["A1"].font = Font(bold=True, size=12)
        for name, sheet in xw.sheets.items():
            for col in sheet.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 48)
            for row in sheet.iter_rows():
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0%" if any(
                            s in str(sheet.cell(row=c.row, column=1).value) + str(sheet.cell(row=1 if name != "Peer comps" else 2, column=c.column).value)
                            for s in ("growth", "margin", "Premium to", "Share of", "Offer vs")
                        ) else "#,##0.00"
        for r in (xw.sheets["Valuation"][1],):
            for c in r:
                c.font, c.fill = bold, head
        for c in ws[2]:
            c.font, c.fill = bold, head

        chart = out_xlsx.with_suffix(".png")
        if chart.exists():
            xw.sheets["Valuation"].add_image(XLImage(str(chart)), "F2")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--peers-csv", help="Use peer figures from a CSV instead of Yahoo Finance")
    ap.add_argument("--out", default="wbd_comps.xlsx", help="Output Excel file")
    ap.add_argument("--as-of", default=pd.Timestamp.today().strftime("%d %b %Y"),
                    help="Date label for the peer data")
    args = ap.parse_args()

    if args.peers_csv:
        raw = pd.read_csv(args.peers_csv)
        print(f"Read {len(raw)} peers from {args.peers_csv}")
    else:
        print("Fetching peer data from Yahoo Finance…")
        raw = fetch_live(PEERS.keys())
    if raw.empty:
        sys.exit("No peer data. Check your internet connection or use --peers-csv.")

    comps = build_comps(raw)
    vals = value_wbd(comps)
    whole_mult = vals.pop("_whole_median_mult", None)
    sotp = vals.pop("_sotp_detail", None)
    ranges = {**vals, **ADVISER_RANGES}
    metrics = deal_metrics(whole_mult)

    out = Path(args.out)
    football_field(ranges, out.with_suffix(".png"), args.as_of)
    write_excel(comps, ranges, metrics, sotp, out, args.as_of)

    pd.set_option("display.width", 140)
    print("\nPeer multiples:")
    print(comps[["ticker", "group", "ev_ebitda", "ev_revenue", "pe"]].round(1).to_string(index=False))
    print("\nImplied WBD value per share vs $31.00 offer:")
    for k, (lo, hi) in ranges.items():
        print(f"  {k:<28} ${lo:6.2f} – ${hi:6.2f}")
    print("\nDeal metrics:")
    for k, v in metrics.items():
        pct = k.startswith(("Premium to", "Share of"))
        print(f"  {k:<52} {v:.1%}" if pct else f"  {k:<52} {v:,.1f}")
    print(f"\nSaved {out} and {out.with_suffix('.png')}")


if __name__ == "__main__":
    main()